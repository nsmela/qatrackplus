"""Browser tests for the reports page.

The reports page is almost entirely JavaScript (reports/js/reports.js): the
filter form for a report type is fetched by ajax when the type is chosen, the
preview is fetched by ajax, saved reports are loaded, deleted and scheduled
through ajax calls and Bootstrap modals, and the date range filters use a
date range picker. A Django test client test cannot exercise any of that,
which is why these tests drive a browser.

Every test waits for the page to show the outcome, and for jQuery to have no
requests in flight, before it looks at the database.
"""

import csv
import os
import shutil
import tempfile
import zipfile
from unittest import mock

import pytest
from django.contrib.auth.models import Group
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from selenium.webdriver.common.by import By

from qatrack.accounts.tests.utils import create_user
from qatrack.faults.tests import utils as fault_utils
from qatrack.qa import models as qa_models
from qatrack.qa.tests import utils as qa_utils
from qatrack.qatrack_core.dates import format_as_date
from qatrack.qatrack_core.tests.live import SeleniumTestCase
from qatrack.reports import faults, filters, models, qc, service_log
from qatrack.reports import reports as reports_module
from qatrack.service_log import models as sl_models
from qatrack.service_log.tests import utils as sl_utils
from qatrack.units.models import Site


def mid_last_month():
    """Noon on the 15th-or-so of last month, which the default "Last Month" work completed filter includes."""
    first_of_month = timezone.localtime().replace(day=1, hour=12, minute=0, second=0, microsecond=0)
    return first_of_month - timezone.timedelta(days=15)


class ReportsSeleniumTestCase(SeleniumTestCase):
    """Common data and helpers for the reports page.

    One site with one unit ("Linac Alpha") and a group ("Therapists") the user
    belongs to. Each test adds what else it needs; `create_qc` assigns and
    completes a one-test list ("Daily Output" with "Output Dose" by default).
    """

    def setUp(self):
        super().setUp()
        qa_utils.create_status(is_default=True)
        self.group = Group.objects.create(name='Therapists')
        self.user.groups.add(self.group)
        # A site needs a slug, or the QC menu on the page shown after logging in fails to render.
        self.site = Site.objects.create(name='Main Site', slug='main-site')
        self.unit = qa_utils.create_unit(name='Linac Alpha', site=self.site)

    def create_qc(self, unit=None, test_list_name='Daily Output', test_name='Output Dose', work_completed=None,
                  value=1.23):
        """Assign a one-test list to `unit` and complete one session of it, with `value` for the test."""
        unit = unit or self.unit
        test = qa_utils.create_test(name=test_name)
        test_list = qa_utils.create_test_list(name=test_list_name)
        qa_utils.create_test_list_membership(test_list=test_list, test=test)
        utc = qa_utils.create_unit_test_collection(unit=unit, test_collection=test_list, assigned_to=self.group)
        tli = qa_utils.create_test_list_instance(
            unit_test_collection=utc, work_completed=work_completed or mid_last_month(), created_by=self.user,
        )
        uti = qa_models.UnitTestInfo.objects.get(unit=unit, test=test)
        qa_utils.create_test_instance(
            test_list_instance=tli, unit_test_info=uti, value=value, created_by=self.user,
            work_completed=tli.work_completed,
        )
        return utc, tli

    def create_saved_report(self, title='Monthly QC', created_by=None, **kwargs):
        user = created_by or self.user
        defaults = {
            'report_type': qc.TestListInstanceSummaryReport.report_type,
            'report_format': 'pdf',
            'title': title,
            'filters': {'work_completed': 'Last Month'},
            'created_by': user,
            'modified_by': user,
        }
        defaults.update(kwargs)
        return models.SavedReport.objects.create(**defaults)

    # -- select2 -----------------------------------------------------------------------------------

    def select2_options(self, el_id):
        """The text of every <option> in <select id=el_id>."""
        return self.driver.execute_script(
            'return Array.from(document.getElementById(arguments[0]).options).map(function(o) {'
            '  return o.text.trim(); });', el_id
        )

    def open_reports(self, query=''):
        """Log in, open the reports page and wait for its saved reports table to load."""
        self.login()
        self.open(reverse('reports') + query)
        self.wait_for_visible('#id_root-report_type + .select2')
        self.wait_for_ajax()
        self.browser_log()  # start the console log afresh for this page

    def choose_report_type(self, name, filter_field):
        """Choose a report type and wait for its filter form (which must contain `filter_field`) to arrive.

        The title box is emptied first, as the page only fills it in with the
        report type's name when it is empty.
        """
        title = self.driver.find_element(By.ID, 'id_root-title')
        title.clear()
        self.select2_search('id_root-report_type', str(name), str(name), exact=True)
        self.wait_for_ajax()
        self.wait_for_element('#filters #%s' % filter_field)
        self.wait.until(lambda d: not self.is_disabled('preview'))

    def preview(self):
        """Click Preview and return the text of the preview once it has been drawn."""
        self.driver.execute_script("document.getElementById('report').innerHTML = '';")
        self.click_css('#preview')
        self.wait_for_ajax()
        self.wait.until(lambda d: d.find_elements(By.CSS_SELECTOR, '#report .container-fluid, .error-message'))
        errors = [e.text for e in self.driver.find_elements(By.CSS_SELECTOR, '.error-message') if e.text.strip()]
        assert not errors, 'The preview was refused: %s' % errors
        return self.driver.find_element(By.CSS_SELECTOR, '#report').text

    def preview_title(self):
        return self.driver.find_element(By.CSS_SELECTOR, '#report h5').text.strip()

    def assert_no_javascript_errors(self, context=''):
        errors = self.javascript_errors()
        assert not errors, 'Console errors %s: %s' % (context, errors)

    def open_saved_report(self, report):
        """Click a saved report's link in the table and wait for the page to load it (and its preview)."""
        self.click_css('#report-id-%d' % report.pk)
        self.wait.until(lambda d: self.value_of('current') == report.title)
        self.wait_for_ajax()

    def choose_date_range(self, el_id, label):
        """Open the date range picker of input `el_id` and click one of its preset ranges."""
        self.click_css('#%s' % el_id)
        self.click_element(self.wait.until(lambda d: self._visible_range(label)))
        self.wait_for_invisible('.daterangepicker .ranges')

    def _visible_range(self, label):
        for item in self.driver.find_elements(By.CSS_SELECTOR, '.daterangepicker .ranges li[data-range-key]'):
            if item.get_attribute('data-range-key') == label and item.is_displayed():
                return item
        return False

    def download(self, format_label, filename):
        """Choose the report format, click Download and return the path of the file the browser saved."""
        if not hasattr(self.driver, 'execute_cdp_cmd'):
            pytest.skip('Saving downloads to a known folder needs Chromium')
        download_dir = tempfile.mkdtemp(prefix='qatrack-reports-download-')
        self.addCleanup(shutil.rmtree, download_dir, ignore_errors=True)
        self.driver.execute_cdp_cmd('Browser.setDownloadBehavior', {'behavior': 'allow', 'downloadPath': download_dir})

        self.select2_choose('id_root-report_format', format_label)
        self.click_css('#download')
        path = os.path.join(download_dir, filename)
        # Chromium writes to "<name>.crdownload" and renames it when the download is complete.
        self.wait_until(
            lambda: os.path.exists(path) and os.path.getsize(path) > 0, timeout=self.timeout * 3,
            message='%s was not downloaded; the folder holds %s' % (filename, os.listdir(download_dir)),
        )
        return path

    def open_schedule_modal(self, report):
        self.click_css('#report-id-%d-schedule' % report.pk)
        self.wait_for_visible('#schedule-modal #id_schedule-users + .select2')
        self.wait_for_ajax()

    def fill_schedule_time_and_date(self, time_label='07:30'):
        """Pick a time of day and add today's date to the recurrence widget."""
        self.select2_choose('id_schedule-time', time_label)
        self.click_css('#schedule-modal .add-date')


class TestEveryReportType(ReportsSeleniumTestCase):
    """Each report type's filter form loads, and its preview shows the data seeded for it."""

    def check_report_types(self, cases):
        """`cases` is a list of (ReportClass, filter field id, setup callable or None, expected texts)."""
        self.open_reports()
        for report_class, filter_field, setup, expected in cases:
            self.choose_report_type(report_class.name, filter_field)
            assert self.value_of('id_root-title') == str(report_class.name)
            if setup:
                setup()
            text = self.preview()
            assert self.preview_title() == str(report_class.name)
            for value in expected:
                assert value in text, '%r missing from the %s preview' % (value, report_class.name)
            self.assert_no_javascript_errors('for %s' % report_class.name)

    def test_every_qc_report_type_previews_seeded_qc(self):
        """Every QC report type loads its filters and previews the seeded unit, test list and test.

        If one failed, choosing that report would leave the page blank or
        show an empty report, with nothing to tell the physicist why.
        """
        utc, tli = self.create_qc()
        qa_models.UnitTestCollection.objects.filter(pk=utc.pk).update(due_date=timezone.now())

        def choose_utc():
            self.select2_choose('id_unit_test_collection', 'Linac Alpha :: Daily Output')

        def choose_test():
            self.select2_choose('id_unit_test_info__test', 'Output Dose')

        self.check_report_types([
            (qc.TestListInstanceSummaryReport, 'id_work_completed', None, ['Linac Alpha', 'Daily Output']),
            (qc.TestListInstanceDetailsReport, 'id_unit_test_collection', choose_utc,
             ['Linac Alpha', 'Daily Output', 'Output Dose', '1.23']),
            (qc.TestInstanceDetailsReport, 'id_unit_test_info__test', choose_test,
             ['Linac Alpha', 'Output Dose', '1.23']),
            (qc.AssignedQCReport, 'id_assigned_to', None, ['Linac Alpha', 'Daily Output', 'Therapists']),
            (qc.AssignedQCDetailsReport, 'id_refs_tols', None, ['Linac Alpha', 'Daily Output', 'Output Dose']),
            (qc.PaperBackupForms, 'id_refs_tols', None, ['Linac Alpha', 'Daily Output', 'Output Dose']),
            (qc.NextDueDatesReport, 'id_due_date', None, ['Linac Alpha', 'Daily Output']),
            (qc.DueAndOverdueQCReport, 'id_active', None, ['Linac Alpha', 'Daily Output']),
        ])

    def test_every_service_log_report_type_previews_seeded_service_events(self):
        """Every service log report type loads its filters and previews the seeded service event or schedule."""
        self.user.first_name, self.user.last_name = 'Pat', 'Physicist'
        self.user.save()
        usa = sl_utils.create_unit_service_area(unit=self.unit, service_area=sl_utils.create_service_area('Gantry'))
        service_type = sl_utils.create_service_type(name='Preventive')
        se = sl_utils.create_service_event(
            unit_service_area=usa,
            service_type=service_type,
            service_status=sl_utils.create_service_event_status(name='Approved', is_default=True),
            user_created_by=self.user,
            datetime_service=timezone.now() - timezone.timedelta(days=3),
            datetime_created=timezone.now() - timezone.timedelta(days=3),
            problem_description='Gantry squeaks at 270 degrees',
            service_time=timezone.timedelta(hours=2),
            lost_time=timezone.timedelta(hours=1),
        )
        sl_utils.create_hours(service_event=se, user=self.user, time=timezone.timedelta(hours=1, minutes=30))
        template = sl_utils.create_service_event_template(
            name='Annual PM', service_type=service_type, service_area=usa.service_area, created_by=self.user,
        )
        schedule = sl_utils.create_service_event_schedule(
            service_event_template=template, unit_service_area=usa, frequency=qa_utils.create_frequency(),
            assigned_to=self.group,
        )
        sl_models.ServiceEventSchedule.objects.filter(pk=schedule.pk).update(due_date=timezone.now())

        self.check_report_types([
            (service_log.ServiceEventSummaryReport, 'id_include_description', None,
             ['Linac Alpha', 'Gantry', 'Preventive', 'Gantry squeaks at 270 degrees']),
            (service_log.ServiceEventDetailsReport, 'id_datetime_service', None,
             ['Linac Alpha', 'Gantry', 'Gantry squeaks at 270 degrees']),
            (service_log.ServiceTimesReport, 'id_service_type', None, ['Linac Alpha', 'Preventive']),
            (service_log.ServiceEventPersonnelSummaryReport, 'id_unit_service_area__service_area', None,
             ['user (Physicist, Pat)', '1:30']),
            (service_log.ScheduledTemplatesReport, 'id_service_event_template', None,
             ['Linac Alpha', 'Gantry', 'Annual PM', 'Therapists']),
            (service_log.NextScheduledServiceEventsDueDatesReport, 'id_due_date', None,
             ['Linac Alpha', 'Annual PM']),
            (service_log.DueAndOverdueServiceEventScheduleReport, 'id_frequency', None,
             ['Linac Alpha', 'Annual PM']),
        ])

    def test_every_fault_report_type_previews_seeded_faults(self):
        """Both fault report types load their filters and preview the seeded fault."""
        fault_type = fault_utils.create_fault_type(code='MLC-INTERLOCK', description='MLC leaf interlock')
        fault_utils.create_fault(
            unit=self.unit, fault_type=fault_type, user=self.user,
            occurred=timezone.now() - timezone.timedelta(days=2),
        )

        self.check_report_types([
            (faults.FaultSummaryReport, 'id_occurred', None, ['Linac Alpha', 'MLC-INTERLOCK']),
            (faults.FaultDetailsReport, 'id_review_status', None, ['Linac Alpha', 'MLC-INTERLOCK']),
        ])


class TestDownloads(ReportsSeleniumTestCase):
    """The Download button returns the report in the chosen format, with the seeded data in it."""

    def setUp(self):
        super().setUp()
        self.utc, self.tli = self.create_qc()

    def prepare_report(self):
        self.open_reports()
        self.choose_report_type(qc.TestListInstanceSummaryReport.name, 'id_work_completed')
        title = self.driver.find_element(By.ID, 'id_root-title')
        title.clear()
        title.send_keys('Linac Alpha monthly QC')

    def expected_row(self):
        return ['Main Site', 'Linac Alpha', 'Daily Output', format_as_date(self.tli.work_completed)]

    def test_csv_download_contains_the_seeded_session(self):
        """Downloading as CSV must give a CSV file listing the session completed last month.

        This is what gets attached to audits; an empty or wrong file would go unnoticed until then.
        """
        self.prepare_report()
        path = self.download('CSV', 'test-list-instance-summary.csv')

        with open(path, newline='') as f:
            rows = list(csv.reader(f))
        assert ['Report Title:', 'Linac Alpha monthly QC'] in rows
        header = rows.index(['Site', 'Unit', 'Test list', 'Due Date', 'Work Completed', 'Pass/Fail Status', 'Link'])
        data = rows[header + 1:]
        assert len(data) == 1
        site, unit, test_list, _due, work_completed = data[0][:5]
        assert [site, unit, test_list, work_completed] == self.expected_row()
        assert data[0][6].endswith(self.tli.get_absolute_url())

    def test_xlsx_download_contains_the_seeded_session(self):
        """Downloading as Excel must give a workbook listing the session completed last month."""
        self.prepare_report()
        path = self.download('Excel', 'test-list-instance-summary.xlsx')

        assert zipfile.is_zipfile(path)
        sheet = load_workbook(path, read_only=True)['Report']
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        assert ['Report Title:', 'Linac Alpha monthly QC'] in [r[:2] for r in rows]
        header = [r[:7] for r in rows].index(
            ['Site', 'Unit', 'Test list', 'Due Date', 'Work Completed', 'Pass/Fail Status', 'Link']
        )
        data = [r for r in rows[header + 1:] if any(c is not None for c in r)]
        assert len(data) == 1
        site, unit, test_list, _due, work_completed = data[0][:5]
        assert [site, unit, test_list, work_completed] == self.expected_row()

    def test_pdf_download_is_a_pdf(self):
        """Downloading as PDF must give a PDF file (rendered server side, so it works headless too)."""
        self.prepare_report()
        path = self.download('PDF', 'test-list-instance-summary.pdf')

        with open(path, 'rb') as f:
            content = f.read()
        assert content.startswith(b'%PDF-')
        assert len(content) > 1000


class TestSavedReports(ReportsSeleniumTestCase):
    """Saving, loading and deleting saved reports."""

    def test_deleting_a_saved_report_through_its_modal_removes_it(self):
        """Confirming the delete modal must delete the report, take it out of the table and reset the form."""
        report = self.create_saved_report(title='Monthly QC')
        keep = self.create_saved_report(title='Weekly QC')
        self.open_reports()
        self.open_saved_report(report)

        self.click_css('#delete-report')
        assert self.wait_for_visible('#delete-modal .modal-title').text == "Delete 'Monthly QC'?"
        self.click_css('#delete')
        self.wait_for_invisible('#delete-modal')
        self.wait_for_text('#form-success', 'Your report was deleted')
        self.wait_for_ajax()

        self.wait.until(lambda d: not d.find_elements(By.ID, 'report-id-%d' % report.pk))
        assert self.driver.find_elements(By.ID, 'report-id-%d' % keep.pk)
        assert list(models.SavedReport.objects.values_list('pk', flat=True)) == [keep.pk]
        assert self.value_of('report_id') == ''
        assert self.value_of('current') == 'New Report'

    def test_cancelling_the_delete_modal_keeps_the_report(self):
        """Choosing "No, Cancel" must close the modal and leave the report alone."""
        report = self.create_saved_report(title='Monthly QC')
        self.open_reports()
        self.open_saved_report(report)

        self.click_css('#delete-report')
        self.click_css('#delete-modal .modal-footer [data-dismiss=modal]')
        self.wait_for_invisible('#delete-modal')
        self.wait_for_ajax()

        assert models.SavedReport.objects.filter(pk=report.pk).exists()
        assert self.driver.find_elements(By.ID, 'report-id-%d' % report.pk)
        assert self.value_of('report_id') == str(report.pk)

    def test_report_shared_by_another_user_is_saved_as_a_new_report(self):
        """A report someone else owns can be loaded but is saved as the viewer's own copy.

        The owner's report must not change: only its creator may edit it.
        """
        owner = create_user(uname='olivia')
        shared = self.create_saved_report(title='Shared QC', created_by=owner)
        shared.visible_to.add(self.group)
        self.open_reports()
        self.open_saved_report(shared)

        save = self.driver.find_element(By.ID, 'save')
        assert save.text.strip() == 'Save as New'
        assert self.value_of('report_id') == ''
        assert not self.driver.find_elements(By.ID, 'report-id-%d-schedule' % shared.pk)

        self.click_css('#save')
        self.wait_for_text('#form-success', 'Your report was saved')
        self.wait_for_ajax()

        copy = models.SavedReport.objects.exclude(pk=shared.pk).get()
        assert (copy.title, copy.report_type, copy.created_by) == (
            'Shared QC', shared.report_type, self.user,
        )
        shared.refresh_from_db()
        assert (shared.created_by, shared.modified_by) == (owner, owner)
        self.wait_for_element('#report-id-%d' % copy.pk)
        assert self.value_of('report_id') == str(copy.pk)
        assert save.text.strip() == 'Save'

    def test_report_id_in_the_url_loads_the_report_and_previews_it(self):
        """Opening /reports/?report_id=N (the "View On Site" link in a report) must load that report.

        Its title, type and filters are filled in and its preview is drawn,
        filtered as it was saved.
        """
        beta = qa_utils.create_unit(name='Linac Beta', site=self.site)
        self.create_qc(test_list_name='Alpha Output')
        self.create_qc(unit=beta, test_list_name='Beta Output', test_name='Beta Dose')
        report = self.create_saved_report(
            title='Alpha audit',
            filters={'work_completed': 'Last Month', 'unit_test_collection__unit': [str(self.unit.pk)]},
        )

        self.open_reports('?report_id=%d' % report.pk)
        self.wait.until(lambda d: self.value_of('current') == 'Alpha audit')
        self.wait.until(lambda d: d.find_elements(By.CSS_SELECTOR, '#report .container-fluid'))
        self.wait_for_ajax()

        assert self.value_of('report_id') == str(report.pk)
        assert self.value_of('id_root-title') == 'Alpha audit'
        assert self.value_of('id_root-report_type') == qc.TestListInstanceSummaryReport.report_type
        assert self.value_of('id_work_completed') == 'Last Month'
        assert self.driver.execute_script(
            "return $('#id_unit_test_collection__unit').val();"
        ) == [str(self.unit.pk)]
        text = self.driver.find_element(By.CSS_SELECTOR, '#report').text
        assert self.preview_title() == 'Alpha audit'
        assert 'Alpha Output' in text
        assert 'Beta Output' not in text


class TestFilterWidgets(ReportsSeleniumTestCase):
    """The date range picker presets and the multi-select filters."""

    def test_date_range_preset_sets_the_range_and_narrows_the_preview(self):
        """Choosing "Last 7 Days" must put that range in the box and keep older sessions out of the preview."""
        now = timezone.now()
        self.create_qc(test_list_name='Recent Check', work_completed=now - timezone.timedelta(days=3))
        self.create_qc(test_list_name='Older Check', test_name='Older Dose',
                       work_completed=now - timezone.timedelta(days=20))
        self.open_reports()
        self.choose_report_type(qc.TestListInstanceSummaryReport.name, 'id_work_completed')
        assert self.value_of('id_work_completed') == 'Last Month'

        self.choose_date_range('id_work_completed', 'Last 7 Days')

        assert self.value_of('id_work_completed') == 'Last 7 Days'
        text = self.preview()
        assert 'Recent Check' in text
        assert 'Older Check' not in text

    def test_last_180_days_preset_starts_180_days_ago(self):
        """The "Last 180 Days" preset must select a range starting 180 days ago.

        The box shows the label and the server works out the dates itself,
        but the picker holds the preset's own dates: reopened, it highlights
        the wrong range and its calendars show the wrong start date.
        """
        self.open_reports()
        self.choose_report_type(qc.TestListInstanceSummaryReport.name, 'id_work_completed')
        self.choose_date_range('id_work_completed', 'Last 180 Days')
        assert self.value_of('id_work_completed') == 'Last 180 Days'

        # Reopen the picker: it marks the preset matching its selected dates as active.
        self.click_css('#id_work_completed')
        active = self.wait_for_visible('.daterangepicker .ranges li.active')
        assert active.get_attribute('data-range-key') == 'Last 180 Days'

        start, expected = self.driver.execute_script(
            "var d = new Date(); d.setDate(d.getDate() - 180);"
            "function ymd(x) { return x.getFullYear() + '-' + ('0' + (x.getMonth() + 1)).slice(-2)"
            "  + '-' + ('0' + x.getDate()).slice(-2); }"
            "return [$('#id_work_completed').data('daterangepicker').startDate.format('YYYY-MM-DD'), ymd(d)];"
        )
        assert start == expected

    def test_unit_filter_narrows_the_preview_to_the_chosen_unit(self):
        """Choosing one unit in the multi-select Unit filter must leave the other unit's QC out of the preview."""
        beta = qa_utils.create_unit(name='Linac Beta', site=self.site)
        self.create_qc(test_list_name='Alpha Output')
        self.create_qc(unit=beta, test_list_name='Beta Output', test_name='Beta Dose')
        self.open_reports()
        self.choose_report_type(qc.TestListInstanceSummaryReport.name, 'id_work_completed')

        text = self.preview()
        assert 'Alpha Output' in text and 'Beta Output' in text

        self.select2_choose('id_unit_test_collection__unit', 'Main Site :: Linac Alpha')
        text = self.preview()
        assert 'Alpha Output' in text
        assert 'Beta Output' not in text
        assert 'Main Site - Linac Alpha' in text  # the filter is listed in the report details


class TestReportFormat(ReportsSeleniumTestCase):
    """The report format select and the fields that only apply to PDFs."""

    def pdf_fields_shown(self):
        return [
            self.driver.execute_script(
                "return $(arguments[0]).closest('.form-group').is(':visible');", '#%s' % el_id
            )
            for el_id in ('id_root-include_logo', 'id_root-paper_size')
        ]

    def test_choosing_csv_hides_the_pdf_only_fields(self):
        """Logo and paper size only apply to PDFs, so choosing CSV hides them and PDF shows them again."""
        self.open_reports()
        self.choose_report_type(qc.TestListInstanceSummaryReport.name, 'id_work_completed')
        assert self.value_of('id_root-report_format') == 'pdf'
        assert self.pdf_fields_shown() == [True, True]

        self.select2_choose('id_root-report_format', 'CSV')
        self.wait.until(lambda d: self.pdf_fields_shown() == [False, False])

        self.select2_choose('id_root-report_format', 'PDF')
        self.wait.until(lambda d: self.pdf_fields_shown() == [True, True])

    def test_report_type_without_pdf_disables_preview(self):
        """A report type that cannot be a PDF can't be previewed (the preview is the PDF's HTML).

        None of the shipped report types lacks PDF, so a CSV-only report type
        is registered for the length of this test.
        """
        patcher = mock.patch.dict(reports_module.REPORT_REGISTRY)
        patcher.start()
        self.addCleanup(patcher.stop)

        class CsvOnlyAssignmentsReport(reports_module.BaseReport):
            report_type = 'test-csv-only-assignments'
            name = 'CSV Only Assignments'
            category = 'Testing'
            filter_class = filters.UnitTestCollectionFilter
            formats = [reports_module.CSV]

        self.open_reports()
        self.choose_report_type(qc.AssignedQCReport.name, 'id_assigned_to')
        assert not self.is_disabled('preview')

        self.driver.find_element(By.ID, 'id_root-title').clear()
        self.select2_search('id_root-report_type', 'CSV Only Assignments', 'CSV Only Assignments', exact=True)
        self.wait_for_ajax()
        self.wait.until(lambda d: self.is_disabled('preview'))
        assert self.select2_options('id_root-report_format') == ['CSV']
        assert not self.is_disabled('download')
        assert self.pdf_fields_shown() == [False, False]

    def test_loading_a_saved_csv_report_hides_the_pdf_only_fields(self):
        """A saved CSV report, once loaded, must hide the logo and paper size fields like choosing CSV does."""
        report = self.create_saved_report(title='CSV export', report_format='csv')
        self.open_reports()
        self.open_saved_report(report)
        assert self.value_of('id_root-report_format') == 'csv'

        assert self.pdf_fields_shown() == [False, False]

    def test_paper_backup_forms_offer_only_pdf(self):
        """Paper backup forms only exist as PDFs, so PDF must be the only format offered."""
        self.open_reports()
        self.choose_report_type(qc.PaperBackupForms.name, 'id_refs_tols')
        assert self.select2_options('id_root-report_format') == ['PDF']
        assert not self.is_disabled('preview')


class TestScheduling(ReportsSeleniumTestCase):
    """Scheduling a saved report to be emailed."""

    def setUp(self):
        super().setUp()
        self.report = self.create_saved_report(title='Weekly QC')
        self.recipient = create_user(uname='rory', is_superuser=False)
        self.recipient.first_name, self.recipient.last_name = 'Rory', 'Radiographer'
        self.recipient.email = 'rory@example.com'
        self.recipient.save()

    def test_schedule_with_user_and_group_recipients_is_saved(self):
        """Choosing a user and a group as recipients must save a schedule sending to both."""
        self.open_reports()
        self.open_schedule_modal(self.report)
        assert self.driver.find_element(By.CSS_SELECTOR, '#schedule-modal .modal-title').text == (
            'Update schedule for Weekly QC'
        )
        self.fill_schedule_time_and_date('07:30')
        self.select2_choose('id_schedule-users', 'rory')
        self.select2_choose('id_schedule-groups', 'Therapists')

        self.click_css('#schedule')
        self.wait_for_text('#schedule-message', 'Schedule updated successfully!')
        self.wait_for_ajax()

        schedule = models.ReportSchedule.objects.get(report=self.report)
        assert list(schedule.users.all()) == [self.recipient]
        assert list(schedule.groups.all()) == [self.group]
        assert schedule.time.strftime('%H:%M') == '07:30'
        assert schedule.emails == ''
        assert '"Rory Radiographer"<rory@example.com>' in schedule.recipients()
        assert timezone.localtime().strftime('%Y%m%d') in str(schedule.schedule)

    def test_schedule_without_recipients_is_refused_with_a_message(self):
        """A schedule with no user, group or email must not be saved, and the modal must say why."""
        self.open_reports()
        self.open_schedule_modal(self.report)
        self.fill_schedule_time_and_date('07:30')

        self.click_css('#schedule')
        self.wait_for_text('#schedule-form .alert-danger', 'You must select at least one group, user, or email address!')
        self.wait_for_ajax()

        assert not models.ReportSchedule.objects.exists()
        assert not self.is_shown('#schedule-message')
